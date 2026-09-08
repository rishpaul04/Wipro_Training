import json
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def create_excel_from_json(json_path, xlsx_path):
    with open(json_path, "r") as f:
        data = json.load(f)

    wb = Workbook()

    # --- Sheet 1: Test Cases ---
    ws = wb.active
    ws.title = "Test Cases"

    headers = ["Test ID", "Description", "Search Product", "Quantity", "Email", "Password", "Expected Result"]
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    for row_idx, tc in enumerate(data["test_cases"], 2):
        values = [
            tc["test_id"],
            tc["description"],
            tc.get("search_product", "N/A"),
            tc.get("quantity", "1"),
            tc.get("email", "N/A"),
            tc.get("password", "N/A"),
            tc["expected_result"],
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center")

    for col in ws.columns:
        max_length = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = max_length + 4

    # --- Sheet 2: Credentials ---
    ws2 = wb.create_sheet("Credentials")
    cred_headers = ["Key", "Value"]
    for col, header in enumerate(cred_headers, 1):
        cell = ws2.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border

    creds = data["credentials"]
    for row_idx, (key, value) in enumerate(creds.items(), 2):
        ws2.cell(row=row_idx, column=1, value=key).border = thin_border
        ws2.cell(row=row_idx, column=2, value=value).border = thin_border

    # --- Sheet 3: URLs ---
    ws3 = wb.create_sheet("URLs")
    url_headers = ["Name", "URL"]
    for col, header in enumerate(url_headers, 1):
        cell = ws3.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border

    urls = data["urls"]
    for row_idx, (key, value) in enumerate(urls.items(), 2):
        ws3.cell(row=row_idx, column=1, value=key).border = thin_border
        ws3.cell(row=row_idx, column=2, value=value).border = thin_border

    wb.save(xlsx_path)
    print(f"[INFO] Excel file created: {xlsx_path}")


if __name__ == "__main__":
    base = os.path.dirname(os.path.abspath(__file__))
    create_excel_from_json(
        os.path.join(base, "test_data.json"),
        os.path.join(base, "test_data.xlsx"),
    )
